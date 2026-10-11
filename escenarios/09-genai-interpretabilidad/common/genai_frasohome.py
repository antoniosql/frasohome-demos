from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import joblib
import numpy as np
import pandas as pd
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


def find_project_root(start: Path | None = None) -> Path:
    current = Path(start or Path.cwd()).resolve()
    for candidate in [current, *current.parents]:
        if (candidate / "datos").exists() and (candidate / "escenarios").exists():
            return candidate
    raise FileNotFoundError("No se ha encontrado la raíz del proyecto FraSoHome.")


def demo_dir(project_root: Path | None = None) -> Path:
    root = find_project_root(project_root)
    return root / "escenarios" / "09-genai-interpretabilidad"


def sales_demo_dir(project_root: Path | None = None) -> Path:
    root = find_project_root(project_root)
    return root / "escenarios" / "04-analitica-ventas-ml"


def load_sales_artifact(project_root: Path | None = None, prefer_students: bool = False) -> dict[str, Any]:
    base = sales_demo_dir(project_root) / "artifacts"
    candidates = []
    if prefer_students:
        candidates.append(base / "alumnos_sales_models_artifact.joblib")
    candidates.append(base / "sales_models_artifact.joblib")
    candidates.append(base / "alumnos_sales_models_artifact.joblib")
    for path in candidates:
        if path.exists():
            artifact = joblib.load(path)
            artifact["_path"] = path
            return artifact
    raise FileNotFoundError(
        "No existe el artefacto de ventas. Ejecuta primero "
        "escenarios/04-analitica-ventas-ml/notebooks_instructor/"
        "01_demo_ventas_feature_engineering_modelos.ipynb"
    )


def standardize_id(series: pd.Series) -> pd.Series:
    out = series.astype(str).str.strip().str.upper().str.replace(r"\s+", "", regex=True)
    return out.replace({"": np.nan, "NAN": np.nan, "NONE": np.nan, "NULL": np.nan})


def load_crm(project_root: Path | None = None) -> pd.DataFrame:
    root = find_project_root(project_root)
    path = root / "datos" / "retail_omnicanal" / "raw_csv" / "crm.csv"
    crm = pd.read_csv(path, dtype=str)
    crm["customer_id_std"] = standardize_id(crm["customer_id"])
    return crm


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            line = " | ".join([c for c in cells if c])
            if line:
                parts.append(line)
    return "\n".join(parts)


def load_knowledge_documents(project_root: Path | None = None) -> pd.DataFrame:
    root = find_project_root(project_root)
    kb = root / "datos" / "rag_copilot_studio" / "knowledge_base" / "Documentos_Knowledge_Clasificados"
    records = []
    for path in sorted(kb.rglob("*.docx")):
        rel = path.relative_to(kb)
        name = path.stem
        records.append(
            {
                "doc_id": name.split("_")[0],
                "file_name": path.name,
                "area": rel.parts[0] if len(rel.parts) > 1 else "general",
                "path": str(path),
                "is_obsolete": "Obsoleta" in path.name or "v1.2" in path.name and "Politica_Devoluciones" in path.name,
                "is_security_test": "Prompt_Injection" in path.name or "NO_PROD" in path.name,
                "text": _docx_text(path),
            }
        )
    return pd.DataFrame(records)


def chunk_documents(documents: pd.DataFrame, max_chars: int = 850, overlap_paragraphs: int = 1) -> pd.DataFrame:
    chunks = []
    for _, doc in documents.iterrows():
        paragraphs = [p.strip() for p in re.split(r"\n+", doc["text"]) if p.strip()]
        buffer: list[str] = []
        chunk_idx = 0
        for paragraph in paragraphs:
            candidate = "\n".join(buffer + [paragraph])
            if len(candidate) > max_chars and buffer:
                text = "\n".join(buffer)
                chunks.append(
                    {
                        "chunk_id": f"{doc['doc_id']}-CH{chunk_idx:03d}",
                        "doc_id": doc["doc_id"],
                        "file_name": doc["file_name"],
                        "area": doc["area"],
                        "is_obsolete": bool(doc["is_obsolete"]),
                        "is_security_test": bool(doc["is_security_test"]),
                        "text": text,
                    }
                )
                chunk_idx += 1
                buffer = buffer[-overlap_paragraphs:] if overlap_paragraphs else []
            buffer.append(paragraph)
        if buffer:
            chunks.append(
                {
                    "chunk_id": f"{doc['doc_id']}-CH{chunk_idx:03d}",
                    "doc_id": doc["doc_id"],
                    "file_name": doc["file_name"],
                    "area": doc["area"],
                    "is_obsolete": bool(doc["is_obsolete"]),
                    "is_security_test": bool(doc["is_security_test"]),
                    "text": "\n".join(buffer),
                }
            )
    return pd.DataFrame(chunks)


@dataclass
class RagIndex:
    chunks: pd.DataFrame
    vectorizer: TfidfVectorizer
    matrix: Any


def build_rag_index(chunks: pd.DataFrame, include_obsolete: bool = False, include_security_test: bool = False) -> RagIndex:
    df = chunks.copy()
    if not include_obsolete:
        df = df[~df["is_obsolete"]].copy()
    if not include_security_test:
        df = df[~df["is_security_test"]].copy()
    vectorizer = TfidfVectorizer(
        lowercase=True,
        strip_accents="unicode",
        ngram_range=(1, 2),
        min_df=1,
        max_df=0.95,
    )
    matrix = vectorizer.fit_transform(df["text"].fillna(""))
    return RagIndex(chunks=df.reset_index(drop=True), vectorizer=vectorizer, matrix=matrix)


def retrieve(index: RagIndex, query: str, top_k: int = 5) -> pd.DataFrame:
    q = index.vectorizer.transform([query])
    sims = cosine_similarity(q, index.matrix).ravel()
    order = np.argsort(sims)[::-1][:top_k]
    out = index.chunks.iloc[order].copy()
    out["score"] = sims[order]
    out["rank"] = range(1, len(out) + 1)
    return out[["rank", "score", "chunk_id", "doc_id", "file_name", "area", "text", "is_obsolete", "is_security_test"]]


def split_expected_terms(expected: str) -> list[str]:
    return [t.strip().lower() for t in re.split(r";|\|", str(expected)) if t.strip()]


def keyword_coverage(text: str, expected: str) -> float:
    terms = split_expected_terms(expected)
    if not terms:
        return np.nan
    clean = normalize_text(text)
    hits = sum(1 for term in terms if normalize_text(term) in clean)
    return hits / len(terms)


def normalize_text(text: str) -> str:
    text = str(text).lower()
    replacements = str.maketrans("áéíóúüñ", "aeiouun")
    text = text.translate(replacements)
    return re.sub(r"\s+", " ", text)


def grounded_answer(query: str, retrieved: pd.DataFrame, max_sentences: int = 4) -> dict[str, Any]:
    query_terms = set(re.findall(r"\w{4,}", normalize_text(query)))
    candidate_sentences = []
    for _, row in retrieved.iterrows():
        sentences = re.split(r"(?<=[.!?])\s+|\n+", row["text"])
        for sentence in sentences:
            sent_norm = normalize_text(sentence)
            overlap = len(query_terms & set(re.findall(r"\w{4,}", sent_norm)))
            if overlap > 0:
                candidate_sentences.append((overlap, row["chunk_id"], row["file_name"], sentence.strip()))
    candidate_sentences = sorted(candidate_sentences, key=lambda x: x[0], reverse=True)
    selected = []
    seen = set()
    for _, chunk_id, file_name, sentence in candidate_sentences:
        key = normalize_text(sentence)[:90]
        if sentence and key not in seen:
            selected.append((chunk_id, file_name, sentence))
            seen.add(key)
        if len(selected) >= max_sentences:
            break
    if not selected and len(retrieved):
        top = retrieved.iloc[0]
        selected = [(top["chunk_id"], top["file_name"], top["text"][:350])]
    answer = " ".join([sentence for _, _, sentence in selected])
    citations = [{"chunk_id": c, "file_name": f} for c, f, _ in selected]
    return {"answer": answer, "citations": citations, "used_chunks": [c for c, _, _ in selected]}


def evaluate_rag_questions(index: RagIndex, test_set: pd.DataFrame, top_k: int = 5) -> pd.DataFrame:
    rows = []
    for _, item in test_set.iterrows():
        query = item["Question"]
        expected = item["Expected response"]
        retrieved = retrieve(index, query, top_k=top_k)
        answer_obj = grounded_answer(query, retrieved)
        retrieved_text = "\n".join(retrieved["text"].tolist())
        answer = answer_obj["answer"]
        rows.append(
            {
                "question": query,
                "expected": expected,
                "top_doc": retrieved.iloc[0]["file_name"] if len(retrieved) else None,
                "top_score": float(retrieved.iloc[0]["score"]) if len(retrieved) else np.nan,
                "retrieval_keyword_coverage": keyword_coverage(retrieved_text, expected),
                "answer_keyword_coverage": keyword_coverage(answer, expected),
                "citation_count": len(answer_obj["citations"]),
                "answer": answer,
            }
        )
    return pd.DataFrame(rows)


def simple_answer_rubric(answer: str, context: str, expected: str = "") -> dict[str, float]:
    answer_norm = normalize_text(answer)
    context_norm = normalize_text(context)
    answer_tokens = set(re.findall(r"\w{4,}", answer_norm))
    context_tokens = set(re.findall(r"\w{4,}", context_norm))
    groundedness = len(answer_tokens & context_tokens) / max(len(answer_tokens), 1)
    return {
        "correctness_keyword_coverage": float(keyword_coverage(answer, expected)) if expected else np.nan,
        "groundedness_token_overlap": float(groundedness),
        "has_caveat_or_limit": float(any(w in answer_norm for w in ["si falta", "no hay evidencia", "escalar", "comprobar", "cautela"])),
        "length_ok": float(60 <= len(answer) <= 900),
    }


def transformed_feature_names(pipeline) -> list[str]:
    return list(pipeline.named_steps["preprocess"].get_feature_names_out())


def local_feature_contributions(result: dict[str, Any], row_index: int | None = None, top_n: int = 12) -> tuple[float, pd.DataFrame, pd.DataFrame]:
    pipeline = result["pipeline"]
    X_test = result["X_test"]
    scores = pipeline.predict_proba(X_test)[:, 1]
    if row_index is None:
        row_index = int(np.argmax(scores))
    row_raw = X_test.iloc[[row_index]].copy()
    row_trans = pipeline.named_steps["preprocess"].transform(row_raw)
    names = transformed_feature_names(pipeline)
    model = pipeline.named_steps["model"]
    if hasattr(model, "feature_importances_"):
        weights = np.asarray(model.feature_importances_)
        values = np.asarray(row_trans).ravel()
        contributions = weights * np.abs(values)
        score_name = "local_weighted_signal"
    elif hasattr(model, "coef_"):
        weights = model.coef_[0]
        values = np.asarray(row_trans).ravel()
        contributions = weights * values
        score_name = "local_linear_contribution"
    else:
        contributions = np.asarray(row_trans).ravel()
        score_name = "local_signal"
    out = pd.DataFrame({"feature": names, score_name: contributions})
    out["abs_value"] = out[score_name].abs()
    out = out.sort_values("abs_value", ascending=False).head(top_n).reset_index(drop=True)
    return float(scores[row_index]), row_raw, out


def summarize_customer_context(row_raw: pd.DataFrame, contributions: pd.DataFrame) -> dict[str, Any]:
    row = row_raw.iloc[0].to_dict()
    return {
        "cliente": row.get("customer_id_std", "cliente seleccionado"),
        "recency_days": row.get("recency_days"),
        "frequency_docs": row.get("frequency_docs"),
        "monetary_gross": row.get("monetary_gross"),
        "tier": row.get("tier_fidelizacion"),
        "last_channel": row.get("last_channel"),
        "last_category": row.get("last_category"),
        "top_signals": contributions["feature"].head(5).tolist(),
    }


def recommendation_from_score(score: float, customer_context: dict[str, Any], evidence: pd.DataFrame) -> dict[str, Any]:
    tier = customer_context.get("tier") or "sin tier informado"
    recency = customer_context.get("recency_days")
    category = customer_context.get("last_category") or "categoría reciente no identificada"
    evidence_files = evidence["file_name"].head(3).tolist() if len(evidence) else []
    action = (
        f"Priorizar una acción de reactivación para cliente {tier}. "
        f"El score de inactividad es {score:.2f}; la recencia observada es {recency} días. "
        f"Personalizar el mensaje con novedades o ventaja moderada relacionada con {category}. "
        "Antes de activar, comprobar consentimiento, presión comercial y margen esperado."
    )
    return {
        "recommendation": action,
        "evidence_files": evidence_files,
        "caveats": [
            "El score es predictivo, no causal.",
            "La recomendación debe respetar consentimiento y política CRM.",
            "Conviene medir uplift o conversión incremental antes de escalar.",
        ],
    }


def build_agent_trace(user_request: str) -> list[dict[str, Any]]:
    return [{"step": 0, "type": "user_request", "detail": user_request}]


def append_trace(trace: list[dict[str, Any]], step_type: str, detail: str, payload: Any = None) -> None:
    trace.append({"step": len(trace), "type": step_type, "detail": detail, "payload": payload})


def evaluate_agent_trace(trace: list[dict[str, Any]]) -> pd.DataFrame:
    types = {item["type"] for item in trace}
    checks = [
        ("comprensión", "intent" in types),
        ("planificación", "plan" in types),
        ("uso de herramientas", "tool_call" in types and "tool_result" in types),
        ("evidencia", "evidence" in types),
        ("seguridad", "guardrail" in types),
        ("respuesta final", "final_answer" in types),
    ]
    return pd.DataFrame(
        {
            "capa": name,
            "cumple": ok,
            "comentario": "traza presente" if ok else "revisar: falta evidencia en la traza",
        }
        for name, ok in checks
    )
